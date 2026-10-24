from pathlib import Path
from io import BytesIO
from hashlib import sha256
from zipfile import ZipFile
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

ROOT = Path(r'D:\IOT-asm')
WORK = ROOT / 'tmp/bai2_2_finish'
SOURCE = ROOT / 'bai2_2/Bao_cao_Bai_2_2_MQTT_Nguyen_Tri_Hieu_202416203.docx'
OUTPUT = SOURCE.with_name(SOURCE.stem + '_Hoan_chinh.docx')
original = SOURCE.read_bytes()
(WORK / 'original.docx').write_bytes(original)
doc = Document(BytesIO(original))
body = doc.element.body
all_ps = list(body.iter(qn('w:p')))
cut = next(p for p in body if p.tag == qn('w:p') and ''.join(p.itertext()).startswith('5 Thực hành'))
start = list(body).index(cut)
for child in list(body)[start:]:
    if child.tag != qn('w:sectPr'):
        body.remove(child)

for p in doc.paragraphs:
    if p.text.startswith('Chương trình Java vừa gửi dữ liệu'):
        p.text = ('Chương trình Java gửi dữ liệu cảm biến mô phỏng lên broker và nhận bản tin MQTT '
                  'để hiển thị đủ các trường JSON. Báo cáo trình bày ba tệp mã nguồn chính, giải thích '
                  'các khối xử lý và minh chứng thực hành về gửi nhận dữ liệu, xử lý JSON lỗi, '
                  'Last Will, retained và tự kết nối lại.')

def para(text='', style=None, size=11.5, keep=False):
    p = doc.add_paragraph(style=style)
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(5)
    pf.line_spacing = 1.05
    pf.keep_with_next = keep
    pf.keep_together = True
    pf.widow_control = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def picture(name, width):
    p = para(size=11)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(WORK / name), width=Cm(width))
    return p

def figure(title, paragraphs, images):
    p = para(title, 'Caption', size=11.5, keep=True)
    p.runs[0].bold = True
    for text in paragraphs:
        para(text, keep=True)
    for i, (name, width) in enumerate(images):
        p = picture(name, width)
        p.paragraph_format.keep_with_next = i < len(images) - 1

h = para('5 Kết quả thực hành và minh chứng', 'Heading 1', size=15, keep=True)
h.paragraph_format.page_break_before = True
para('Chương trình sử dụng MSSV 202416203 và broker tcp://broker.hivemq.com:1883. '
     'Java đăng ký it4735e/202416203/#; MQTTX được dùng để theo dõi các topic telemetry và status. '
     'Lệnh chạy với các thư viện đã có trong thư mục lib:')
p = para('cd D:\\IOT-asm\\bai2_2\n'
         'java -cp "out;lib/*" vn.edu.iot.MqttAssignment --mssv 202416203', size=9)
for r in p.runs:
    r.font.name = 'Consolas'

figure('Hình 1 Java kết nối và đăng ký topic', [
    'Terminal hiển thị Client ID riêng, kết nối thành công tới HiveMQ và đăng ký '
    'it4735e/202416203/# với QoS 1. Java gửi trạng thái online có retained và bắt đầu '
    'publish bản tin telemetry đầu tiên với packet_no=126.'
], [('image1.png', 17.0)])

figure('Hình 2 Java gửi và nhận telemetry định kỳ', [
    'Các bản tin trên it4735e/202416203/telemetry có packet_no tăng từ 143 đến 148. '
    'Dòng QoS 1 confirmed cho biết broker đã xác nhận bản tin Java gửi. '
    'Do đăng ký topic có ký tự #, Java cũng nhận lại telemetry của chính mình và hiển thị '
    'đủ id, name, packet_no, temperature, humidity, tds và pH.'
], [('image2.png', 16.6)])

figure('Hình 3 MQTTX nhận dữ liệu JSON do Java gửi', [
    'MQTTX nhận các bản tin JSON trên it4735e/202416203/telemetry với packet_no '
    '150, 151 và 152. Payload có name=my-java-mqtt-client và đủ bảy trường dữ liệu cảm biến.',
    'Ảnh cho thấy subscription của MQTTX sử dụng QoS 0, nên broker chuyển bản tin tới '
    'MQTTX ở QoS 0. Java vẫn publish tới broker bằng QoS 1 như log xác nhận ở Hình 2; '
    'hai chặng truyền có thể dùng mức QoS khác nhau.'
], [('image3.png', 17.4)])

figure('Hình 4 Xử lý bản tin sai JSON và tiếp tục hoạt động', [
    'MQTTX publish chuỗi hello dưới dạng Plaintext vào it4735e/202416203/telemetry '
    'với QoS 1. Cảnh báo Received message SyntaxError trên MQTTX xuất hiện do công cụ '
    'đang hiển thị bản tin nhận ở chế độ JSON.',
    'Trong terminal Java, payload này gây thông báo JSON error kèm continuing. '
    'Ngay sau đó, các bản tin packet_no=268 và 269 tiếp tục được gửi, xác nhận QoS 1 '
    'và nhận lại bình thường. Lỗi JSON được xử lý mà không phải khởi động lại chương trình.'
], [('image4.png', 17.0), ('image5.png', 17.0)])

figure('Hình 5 Last Will và trạng thái retained', [
    'MQTTX theo dõi it4735e/202416203/status khi Java đang online. Sau khi tiến trình '
    'Java bị kết thúc bất thường bằng Task Manager, broker phát hiện mất kết nối và '
    'gửi Last Will có payload offline lên topic trạng thái.',
    'Khi đăng ký lại topic status, MQTTX nhận ngay offline với nhãn Retained. '
    'Broker đã lưu trạng thái cuối cùng và gửi lại cho subscription mới. '
    'Tin online được Java chủ động gửi lúc kết nối; Last Will được cấu hình là offline.'
], [('image6.png', 17.0)])

figure('Hình 6 Tự kết nối lại và đăng ký lại sau khi mất mạng', [
    'Khi mạng bị ngắt tạm thời, Java báo lỗi xác nhận gửi, Connection lost và tạm dừng '
    'publish. Sau khi mạng được khôi phục, chương trình tự kết nối lại tới HiveMQ, '
    'hiển thị Reconnected và đăng ký lại it4735e/202416203/# với QoS 1.',
    'Java gửi lại trạng thái online có retained và tiếp tục gửi, nhận telemetry với '
    'packet_no=138, 139 và 140. Các bước phục hồi được thực hiện trong cùng phiên '
    'chạy, không cần khởi động lại chương trình.'
], [('image7.png', 17.0)])

para('Nhận xét sau khi thực hành: ........................................................................................')

doc.core_properties.subject = 'MSSV 202416203 Mã nguồn và minh chứng thực hành MQTT'
doc.save(OUTPUT)

# Confirm all source-code paragraphs remain exactly as supplied.
before = Document(BytesIO(original))
old_code = [p.text for p in before.paragraphs if p.style.name == 'Code']
new_code = [p.text for p in doc.paragraphs if p.style.name == 'Code']
assert old_code == new_code
with ZipFile(BytesIO(original)) as a, ZipFile(OUTPUT) as b:
    a_images = sorted(sha256(a.read(n)).hexdigest() for n in a.namelist() if n.startswith('word/media/') and not n.endswith('/'))
    b_images = sorted(sha256(b.read(n)).hexdigest() for n in b.namelist() if n.startswith('word/media/') and not n.endswith('/'))
    assert a_images == b_images
assert not any(p.text.startswith(('Thực hiện:', 'Cần thấy trong ảnh:')) for p in doc.paragraphs)
assert sum(p.text.startswith('Hình ') for p in doc.paragraphs) == 6
(WORK/'checks.json').write_text(json.dumps({'source_code_preserved': True, 'images_preserved': len(b_images), 'figures':6, 'reflection':'left blank', 'output':str(OUTPUT)}, ensure_ascii=False, indent=2), encoding='utf-8')
print(OUTPUT)
