from pathlib import Path
from io import BytesIO
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile
import json
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_LINE_SPACING
from docx.oxml.ns import qn

ROOT = Path(r'D:\IOT-asm')
WORK = ROOT/'tmp/bai2_2_compact'
WORK.mkdir(exist_ok=True)
SOURCE = ROOT/'bai2_2/Bao_cao_Bai_2_2_MQTT_Nguyen_Tri_Hieu_202416203_Hoan_chinh.docx'
raw = SOURCE.read_bytes()
(WORK/'before.docx').write_bytes(raw)
doc = Document(BytesIO(raw))
body = doc.element.body
txt = lambda el: ''.join(el.xpath('.//w:t/text()'))
start = next(i for i,el in enumerate(body) if el.tag==qn('w:p') and txt(el)=='2 Cấu hình kết nối')
end = next(i for i,el in enumerate(body) if el.tag==qn('w:p') and txt(el).startswith('5 Kết quả thực hành'))
anchor = body[end]
for el in list(body)[start:end]:
    body.remove(el)

name='Ma nguon day du'
style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
style.base_style=doc.styles['Normal']
style.font.name='Consolas'
style.font.size=Pt(9)
style.font.color.rgb=RGBColor(0,0,0)
fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
for key in ('ascii','hAnsi','eastAsia','cs'):
    fonts.set(qn('w:'+key),'Consolas')
for key in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
    fonts.attrib.pop(qn('w:'+key),None)
pf=style.paragraph_format
pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY
pf.line_spacing=Pt(10.2)
pf.space_before=Pt(0)
pf.space_after=Pt(0)
pf.keep_with_next=False
pf.keep_together=True
pf.widow_control=False
pf.left_indent=Pt(0)
pf.right_indent=Pt(0)

def insert(p):
    anchor.addprevious(p._p)
    return p

def heading(text, first=False):
    p=doc.add_paragraph(text,'Heading 1')
    p.paragraph_format.page_break_before=first
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(7 if not first else 0)
    p.paragraph_format.space_after=Pt(4)
    for r in p.runs:
        r.font.size=Pt(14)
    return insert(p)

def prose(label, text):
    p=doc.add_paragraph()
    pf=p.paragraph_format
    pf.space_before=Pt(0)
    pf.space_after=Pt(3)
    pf.line_spacing=1.03
    pf.keep_with_next=False
    pf.keep_together=True
    pf.widow_control=True
    for t,bold in [(label,True),(': '+text,False)]:
        r=p.add_run(t)
        r.bold=bold
        r.font.name='Times New Roman'
        r.font.size=Pt(11)
    return insert(p)

sections=[
 ('2 Cấu hình kết nối','ClientConfig.java',[
  ('ClientConfig(...)','Lưu MSSV, broker và khoảng chờ gửi; kiểm tra MSSV hợp lệ, khoảng chờ dương và địa chỉ dạng tcp://host:port.'),
  ('parse(args)','Đọc --mssv, --broker và --interval; dùng HiveMQ và 5 giây làm mặc định, đồng thời báo lỗi tùy chọn sai hoặc thiếu giá trị.'),
  ('telemetryTopic(), subscriptionTopic(), statusTopic()','Tạo lần lượt topic dữ liệu, topic nhận có ký tự # và topic trạng thái theo MSSV.')
 ]),
 ('3 Dữ liệu cảm biến','Telemetry.java',[
  ('sample(packetNo)','Tạo dữ liệu mô phỏng với đủ bảy trường; số thứ tự lấy từ bộ đếm gửi.'),
  ('toJson()','Đổi bản tin thành JSON, giữ đúng tên khóa packet_no và pH.'),
  ('fromJson(payload)','Đọc đủ các trường JSON, kiểm tra nội dung thừa và phát sinh JSONException khi dữ liệu không hợp lệ.'),
  ('display()','Định dạng dữ liệu để in ra terminal; các giá trị cảm biến có một chữ số thập phân.')
 ]),
 ('4 Chương trình gửi nhận MQTT','MqttAssignment.java',[
  ('MqttAssignment(config)','Tạo client ID bằng UUID và đăng ký callback. connectComplete() gọi prepareSession(); connectionLost() tạm dừng gửi; messageArrived() giải mã UTF-8 rồi gọi formatReceived(). deliveryComplete() để trống vì worker tự chờ xác nhận gửi.'),
  ('connectionOptions()','Cấu hình MQTT 3.1.1, clean session, tự kết nối lại và Last Will offline có retained trên topic status, QoS 1.'),
  ('run()','Thử kết nối lần đầu; nếu lỗi thì chờ 5 giây rồi thử lại. Khi kết nối được, lên lịch publishTelemetry() và giữ chương trình hoạt động.'),
  ('prepareSession()','Sau mỗi lần kết nối, đăng ký lại topic # với QoS 1, gửi online có retained rồi cho phép gửi telemetry. Lỗi thiết lập phiên được thử lại sau 2 giây.'),
  ('publishTelemetry()','Tạo JSON và publish QoS 1 khi kết nối sẵn sàng; tăng packet_no từ 126 khi lệnh gửi được chấp nhận, sau đó chờ broker xác nhận.'),
  ('formatReceived()','In status dưới dạng văn bản; các bản tin khác được đọc thành Telemetry. Bắt JSONException để thông báo lỗi và tiếp tục nhận bản tin.'),
  ('close()','Dừng worker, chủ động gửi offline có retained khi thoát bình thường, ngắt kết nối và giải phóng tài nguyên một lần.'),
  ('main(args), usage()','Đọc tham số, xử lý --help và lỗi cấu hình; cài shutdown hook gọi close() khi JVM kết thúc. usage() in cú pháp chạy.')
 ])]

manifest=[]
for idx,(title,filename,explanations) in enumerate(sections):
    heading(title,first=idx==0)
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.keep_with_next=True
    r=p.add_run('Tệp '+filename)
    r.bold=True
    r.font.size=Pt(11)
    insert(p)
    lines=(ROOT/'bai2_2/src/main/java/vn/edu/iot'/filename).read_text(encoding='utf-8').splitlines()
    paragraphs=[]
    for line in lines:
        p=doc.add_paragraph(line, name)
        if line.strip().startswith('@'):
            p.paragraph_format.keep_with_next=True
        insert(p)
        paragraphs.append(p)
    p=doc.add_paragraph('Giải thích các hàm','Heading 2')
    p.paragraph_format.space_before=Pt(5)
    p.paragraph_format.space_after=Pt(3)
    p.paragraph_format.keep_with_next=True
    for r in p.runs:
        r.font.size=Pt(11.5)
    insert(p)
    for j,(label,text) in enumerate(explanations):
        p=prose(label,text)
        if len(explanations)<=4 and j<len(explanations)-1:
            p.paragraph_format.keep_with_next=True
    assert [p.text for p in paragraphs]==lines
    manifest.append({'file':filename,'source_lines':len(lines),'explanation_paragraphs':len(explanations)})

OUTPUT=WORK/'compact.docx'
doc.save(OUTPUT)

# Content outside the three-file section and all embedded images remain unchanged.
before=Document(BytesIO(raw))
def body_part(d,mode):
    children=list(d.element.body)
    if mode=='intro':
        stop=next(i for i,el in enumerate(children) if el.tag==qn('w:p') and txt(el)=='2 Cấu hình kết nối')
        return children[:stop]
    begin=next(i for i,el in enumerate(children) if el.tag==qn('w:p') and txt(el).startswith('5 Kết quả thực hành'))
    return children[begin:]
from lxml import etree
for mode in ('intro','photos'):
    a=body_part(before,mode)
    b=body_part(doc,mode)
    assert [etree.tostring(el) for el in a]==[etree.tostring(el) for el in b],mode
with ZipFile(BytesIO(raw)) as a,ZipFile(OUTPUT) as b:
    images=lambda z:sorted(sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/'))
    assert images(a)==images(b)
    assert b.testzip() is None
(WORK/'manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_hash':sha256(raw).hexdigest(),'candidate':str(OUTPUT),'files':manifest,'all_source_lines_preserved':True,'intro_and_photo_section_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUTPUT)
print(manifest)
