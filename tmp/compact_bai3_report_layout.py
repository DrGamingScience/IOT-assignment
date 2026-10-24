from pathlib import Path
from tempfile import NamedTemporaryFile
import shutil

from docx import Document

path = Path(r"D:\IOT-asm\bai3\bai3_1\Bao_cao_Bai3_Phan1_Nhom16.docx")
doc = Document(path)

diagram = next((p for p in doc.paragraphs if p.text.startswith("Producer Python  -->  default exchange")), None)
tree = next((p for p in doc.paragraphs if p.text.startswith("bai3_1/")), None)
assert diagram is not None and tree is not None

diagram.text = (
    "Producer --> default exchange --> it4735.nhom16.telemetry --> Consumer\n"
    "                 routing key = queue name                 <-- ACK"
)
tree.text = (
    "bai3_1/\n"
    "  common.py                 # cau hinh RabbitMQ va queue\n"
    "  producer.py, consumer.py  # gui va nhan thong diep\n"
    "  compose.yaml, Dockerfile  # RabbitMQ va cac container\n"
    "  requirements.txt          # thu vien pika\n"
    "  .env.example, .dockerignore # cau hinh mau va loai file"
)

with NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as tmp:
    temp = Path(tmp.name)
try:
    doc.save(temp)
    check = Document(temp)
    assert len(check.inline_shapes) == 4, "Embedded screenshots changed"
    assert "202400057" in " ".join(c.text for t in check.tables for row in t.rows for c in row.cells)
    shutil.move(str(temp), str(path))
finally:
    temp.unlink(missing_ok=True)

print("Compacted cover-page diagram and file tree; all four screenshots and roster remain.")
