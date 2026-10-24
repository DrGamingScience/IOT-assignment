from pathlib import Path
from io import BytesIO
from copy import deepcopy
from hashlib import sha256
from zipfile import ZipFile
import json
from docx import Document
from docx.oxml.ns import qn

WORK = Path(r'D:\IOT-asm\tmp\bai2_2_hinh3')
SOURCE = Path(r'D:\IOT-asm\bai2_2\Bao_cao_Bai_2_2_MQTT_Nguyen_Tri_Hieu_202416203_Hoan_chinh.docx')
original = SOURCE.read_bytes()
(WORK/'before.docx').write_bytes(original)
doc = Document(BytesIO(original))
paras = doc.paragraphs
start = next(i for i,p in enumerate(paras) if p.text.startswith('Hình 3 '))
assert paras[start+1].text.startswith('MQTTX nhận các bản tin JSON')
assert paras[start+2].text.startswith('Ảnh cho thấy subscription')

replacements = [
    'Hình 3 Java nhận và phân tích JSON gửi từ MQTTX',
    'MQTTX publish bản tin JSON lên it4735e/202416203/telemetry với QoS 1. '
    'Payload kiểm thử gồm id=11, name=mqttx, packet_no=1, temperature=25, '
    'humidity=55, tds=900 và pH=6.5. Thẻ Published màu xanh trong MQTTX '
    'hiển thị bản tin đã gửi.',
    'Terminal Java hiển thị Received trên cùng topic và đọc đủ bảy trường với '
    'các giá trị trùng khớp payload, trong đó name=mqttx và packet_no=1 '
    'phân biệt bản tin kiểm thử với telemetry do Java tự gửi. '
    'Kết quả xác nhận chương trình nhận bản tin từ client MQTTX qua HiveMQ '
    'và phân tích JSON để hiển thị dữ liệu.'
]

for p,text in zip(paras[start:start+3], replacements):
    formatting = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
    p.clear()
    r = p.add_run(text)
    if formatting is not None:
        r._r.insert(0, formatting)

OUTPUT = WORK/'checked.docx'
doc.save(OUTPUT)

# Verify the requested text edit preserves every image and all other text.
old = Document(BytesIO(original))
assert len(old.paragraphs) == len(doc.paragraphs)
for i,(a,b) in enumerate(zip(old.paragraphs,doc.paragraphs)):
    if i not in range(start,start+3):
        assert a.text == b.text, i
with ZipFile(BytesIO(original)) as a, ZipFile(OUTPUT) as b:
    images=lambda z: sorted(sha256(z.read(n)).hexdigest() for n in z.namelist() if n.startswith('word/media/') and not n.endswith('/'))
    assert images(a)==images(b)
    assert b.testzip() is None
assert doc.paragraphs[-1].text.startswith('Nhận xét sau khi thực hành: ...')
(WORK/'edit_manifest.json').write_text(json.dumps({'source':str(SOURCE),'source_hash':sha256(original).hexdigest(),'candidate':str(OUTPUT),'figure3_text':replacements,'all_other_text_and_images_preserved':True},ensure_ascii=False,indent=2),encoding='utf-8')
print(OUTPUT)
