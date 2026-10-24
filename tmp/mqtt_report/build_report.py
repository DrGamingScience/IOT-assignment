from pathlib import Path
import ast

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\IOT-asm\bai2_1")
OUTPUT = ROOT / "Bao_cao_Bai_2_MQTT_Nhom_16_Royal_Things.docx"
sources = {name: (ROOT / name).read_text(encoding="utf-8")
           for name in ("publisher.py", "subscriber.py")}
topic = next(ast.literal_eval(node.value)
             for node in ast.parse(sources["publisher.py"]).body
             if isinstance(node, ast.Assign)
             and any(isinstance(target, ast.Name) and target.id == "TOPIC"
                     for target in node.targets))

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.65)
section.left_margin = section.right_margin = Inches(0.75)
section.footer_distance = Inches(0.25)

for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
    style = doc.styles[name]
    style.font.name = "Times New Roman"
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.underline = False
    style.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.paragraph_format.space_after = Pt(6)
    ppr = style.element.find(qn("w:pPr"))
    if ppr is not None:
        for border in list(ppr.findall(qn("w:pBdr"))):
            ppr.remove(border)
doc.styles["Normal"].font.size = Pt(12)
doc.styles["Normal"].paragraph_format.line_spacing = 1.12
doc.styles["Title"].font.size = Pt(22)
doc.styles["Title"].font.bold = True
doc.styles["Subtitle"].font.size = Pt(15)
doc.styles["Heading 1"].font.size = Pt(15)
doc.styles["Heading 1"].font.bold = True
doc.styles["Heading 1"].paragraph_format.space_before = Pt(8)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(8)
doc.styles["Caption"].font.size = Pt(11)
doc.styles["Caption"].font.italic = True


def table_style(table):
    table.autofit = False
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement(f"w:{side}")
        for key, value in {"val": "single", "sz": "6", "color": "D9D9D9"}.items():
            edge.set(qn(f"w:{key}"), value)
        borders.append(edge)
    props.append(borders)
    margins = OxmlElement("w:tblCellMar")
    for side, value in (("top", "95"), ("bottom", "95"), ("left", "110"), ("right", "110")):
        item = OxmlElement(f"w:{side}")
        item.set(qn("w:w"), value)
        item.set(qn("w:type"), "dxa")
        margins.append(item)
    props.append(margins)
    for row in table.rows:
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05


def code_block(source):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(10.5)
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.widow_control = False
    run = paragraph.add_run(source.rstrip("\n"))
    run.font.name = "Consolas"
    run.font.size = Pt(9.5)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Consolas")


def screenshot_frame(number, name, description):
    caption = doc.add_paragraph(f"Hình {number}  {name}", style="Caption")
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True
    table = doc.add_table(rows=1, cols=1)
    table.columns[0].width = Inches(7)
    table.rows[0].height = Inches(2.05)
    table.rows[0].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    table_style(table)
    cell = table.cell(0, 0)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("CHÈN ẢNH DEMO TẠI ĐÂY")
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string("777777")
    run = p.add_run("\n" + description)
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string("777777")
    gap = doc.add_paragraph()
    gap.paragraph_format.space_after = Pt(1)
    gap.paragraph_format.space_before = Pt(0)
    gap.paragraph_format.line_spacing = Pt(2)
    gap.add_run().font.size = Pt(2)


title = doc.add_paragraph("BÁO CÁO THỰC HÀNH MQTT", style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(16)
title.paragraph_format.space_after = Pt(10)
subtitle = doc.add_paragraph("Bài 2 Trao đổi dữ liệu sử dụng giao thức MQTT", style="Subtitle")
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
group = doc.add_paragraph("Nhóm 16 Royal Things")
group.alignment = WD_ALIGN_PARAGRAPH.CENTER
group.runs[0].bold = True
group.runs[0].font.size = Pt(16)
group.paragraph_format.space_after = Pt(20)

doc.add_heading("Thành viên nhóm", level=1)
members = [
    ("1", "Nguyễn Đăng Long", "202400057", "Thành viên"),
    ("2", "Nguyễn Hoàng Gia", "202400040", "Nhóm trưởng"),
    ("3", "Nguyễn Trí Hiếu", "202416203", "Thành viên"),
    ("4", "Phạm Gia Linh", "202416262", "Thành viên"),
]
table = doc.add_table(rows=1, cols=4)
widths = (0.55, 3.0, 1.75, 1.7)
for column, width in zip(table.columns, widths):
    column.width = Inches(width)
for cell, text in zip(table.rows[0].cells, ("STT", "Họ và tên", "MSSV", "Vai trò")):
    cell.text = text
    shade = OxmlElement("w:shd")
    shade.set(qn("w:fill"), "E7E6E6")
    cell._tc.get_or_add_tcPr().append(shade)
    cell.paragraphs[0].runs[0].bold = True
repeat_header = OxmlElement("w:tblHeader")
table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
for member in members:
    for cell, text in zip(table.add_row().cells, member):
        cell.text = text
table_style(table)
for row in table.rows:
    for cell, width in zip(row.cells, widths):
        cell.width = Inches(width)
    for index in (0, 2):
        row.cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph()
doc.add_heading("Nội dung thực hiện", level=1)
doc.add_paragraph(
    "Chương trình publisher gửi dữ liệu nhiệt độ và độ ẩm dạng JSON mỗi 5 giây. "
    "Chương trình subscriber nhận bản tin từ broker, phân tích JSON và hiển thị dữ liệu. "
    "MQTTX Web được dùng để kiểm tra bản tin do broker chuyển tiếp."
)
doc.add_paragraph("Broker sử dụng: broker.hivemq.com, cổng 1883.")
doc.add_paragraph(f"Topic trong mã nguồn: {topic}.")

doc.add_page_break()
doc.add_heading("Mã nguồn publisher", level=1)
doc.add_paragraph(
    "Publisher chờ kết nối thành công, tạo dữ liệu mô phỏng, chuyển thành JSON và gửi lên topic của nhóm. "
    "Sau mỗi lần gửi, chương trình chờ 5 giây rồi lặp lại."
)
code_block(sources["publisher.py"])

doc.add_page_break()
doc.add_heading("Mã nguồn subscriber", level=1)
doc.add_paragraph(
    "Subscriber đăng ký topic khi kết nối thành công. Khi có bản tin, chương trình đọc JSON và hiển thị "
    "nhiệt độ, độ ẩm cùng thời gian nhận. Bản tin sai định dạng được báo lỗi và bỏ qua."
)
code_block(sources["subscriber.py"])

doc.add_page_break()
doc.add_heading("Ảnh chạy demo", level=1)
doc.add_paragraph("Chèn ảnh chụp màn hình chạy thực tế vào các khung tương ứng.")
screenshot_frame(1, "Publisher gửi dữ liệu", "Terminal hiển thị các lần PUBLISHED cách nhau khoảng 5 giây")
screenshot_frame(2, "Subscriber nhận dữ liệu", "Terminal hiển thị topic cùng nhiệt độ và độ ẩm nhận được")
screenshot_frame(3, "MQTTX Web kiểm tra dữ liệu", f"Topic {topic} và các bản tin JSON nhận được")

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run("Trang ").font.size = Pt(10)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)

doc.core_properties.title = "Báo cáo bài 2 MQTT Nhóm 16 Royal Things"
doc.core_properties.author = "Nhóm 16 Royal Things"
doc.core_properties.subject = "Mã nguồn và ảnh demo bài thực hành MQTT"
doc.save(OUTPUT)
print(OUTPUT)
for name, source in sources.items():
    print(name, "lines:", len(source.splitlines()), "longest line:", max(map(len, source.splitlines())))
