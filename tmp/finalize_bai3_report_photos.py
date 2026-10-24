"""Clarify final report captions to match the group's inserted screenshots."""
from pathlib import Path
from tempfile import NamedTemporaryFile
from docx import Document
import shutil

path = Path(r'D:\IOT-asm\bai3\bai3_1\Bao_cao_Bai3_Phan1_Nhom16.docx')
doc = Document(path)
replacements = {
    'Kết quả kiểm tra cho thấy queue giữ bản tin trong thời gian consumer dừng và chuyển tiếp các bản tin đang chờ khi consumer hoạt động trở lại. Các vị trí chèn ảnh thực nghiệm được bố trí ở cuối báo cáo.':
    'Kết quả kiểm tra cho thấy queue giữ bản tin trong thời gian consumer dừng và chuyển tiếp bản tin đang chờ khi consumer chạy lại. Báo cáo có bốn ảnh chụp thực nghiệm: chương trình chạy, queue bình thường, queue tích bản tin và queue sau khi nhận hết.',
    'Kiểm tra bổ sung gửi chuỗi hello rồi JSON hợp lệ cho kết quả PASS: consumer báo lỗi bản tin đầu và tiếp tục hiển thị temperature=23 C, humidity=33 %. Các ảnh màn hình được chèn tại bốn khung ở các trang sau.':
    'Kiểm tra bổ sung gửi chuỗi hello rồi JSON hợp lệ cho kết quả PASS: consumer báo lỗi bản tin đầu và tiếp tục hiển thị temperature=23 C, humidity=33 %. Bốn ảnh chụp thực tế được trình bày ở hai trang sau.',
    'Thay từng hình giữ chỗ bằng ảnh chụp thực tế: nhấn phải vào hình, chọn Change Picture hoặc Thay đổi ảnh, chọn ảnh đã lưu trong anh_bao_cao. Điền danh sách thành viên ở trang đầu trước khi nộp.':
    'Bốn ảnh chụp thực tế đã được chèn vào báo cáo. Danh sách thành viên nhóm cần được điền ở trang đầu trước khi nộp.',
    'Gợi ý chụp ảnh: giữ được tên queue và các nhãn Ready, Unacked, Consumers trong cùng ảnh. Với chu kỳ 5 giây, biểu đồ tốc độ có thể tạm bằng 0 giữa hai bản tin.':
    'Ảnh chụp giữ được tên queue và các số Ready, Unacked, Consumers. Với chu kỳ 5 giây, biểu đồ tốc độ có thể tạm bằng 0 giữa hai bản tin.',
    'Hình 3  Bản tin chờ trong queue khi consumer dừng. Lần kiểm tra tự động ghi nhận Ready = 6.':
    'Hình 3  Queue tích bản tin khi consumer dừng; ảnh chụp thực tế ghi nhận Ready = 8, Consumers = 0.',
    'Hình 4  Consumer nhận các bản tin tồn đọng; queue trở về trạng thái không còn bản tin chờ.':
    'Hình 4  Consumer nhận các bản tin tồn đọng; Ready = 0, Unacked = 0, Consumers = 1.',
}
found = set()
for paragraph in doc.paragraphs:
    old = paragraph.text
    if old not in replacements:
        continue
    new = replacements[old]
    if paragraph.runs:
        paragraph.runs[0].text = new
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(new)
    found.add(old)
missing = set(replacements) - found
if missing:
    raise RuntimeError(f'Report paragraphs not found: {len(missing)}')
with NamedTemporaryFile(dir=path.parent, suffix='.docx', delete=False) as tmp:
    temp = Path(tmp.name)
try:
    doc.save(temp)
    check = Document(temp)
    assert len(check.inline_shapes) == 4, f'Expected four embedded screenshots, found {len(check.inline_shapes)}'
    assert 'Ready = 8' in '\n'.join(p.text for p in check.paragraphs)
    shutil.move(str(temp), str(path))
finally:
    temp.unlink(missing_ok=True)
print('Updated report captions and verified four embedded screenshots.')
