"""Add concise source archive guide to the report; preserve embedded photos."""
from pathlib import Path
from tempfile import NamedTemporaryFile
import shutil

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

path = Path(r'D:\IOT-asm\bai3\bai3_1\Bao_cao_Bai3_Phan1_Nhom16.docx')
doc = Document(path)
old = next((p for p in doc.paragraphs if p.text.startswith('Các tệp chính gồm common.py,')), None)
assert old is not None, 'Could not find source overview paragraph'
assert not any(p.text == 'Cấu trúc gói mã nguồn' for p in doc.paragraphs)

old.runs[0].text = ('Cảm biến trong bài là dữ liệu mô phỏng, không sử dụng phần cứng thật. '
                    'Báo cáo Word được nộp riêng; ZIP đi kèm chỉ chứa mã nguồn và cấu hình Docker.')
for run in old.runs[1:]:
    run.text = ''

def insert_after(reference, style_name, text):
    element = OxmlElement('w:p')
    reference._p.addnext(element)
    paragraph = Paragraph(element, reference._parent)
    paragraph.style = doc.styles[style_name]
    paragraph.add_run(text)
    return paragraph

heading = insert_after(old, 'Heading 2', 'Cấu trúc gói mã nguồn')
code = insert_after(heading, 'SourceCode',
    'bai3_1/\n'
    '  common.py                  # cau hinh dung chung\n'
    '  producer.py                # gui JSON moi 5 giay\n'
    '  consumer.py                # nhan, hien thi va ACK\n'
    '  compose.yaml, Dockerfile   # RabbitMQ va cac container\n'
    '  requirements.txt           # thu vien pika\n'
    '  .env.example               # mau cau hinh, sao chep thanh .env\n'
    '  .dockerignore')
ppr = code._p.get_or_add_pPr()
shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F4F4F4'); ppr.append(shade)

with NamedTemporaryFile(dir=path.parent, suffix='.docx', delete=False) as tmp:
    temp = Path(tmp.name)
try:
    doc.save(temp)
    check = Document(temp)
    assert len(check.inline_shapes) == 4, 'Embedded screenshots changed unexpectedly'
    assert any(p.text == 'Cấu trúc gói mã nguồn' for p in check.paragraphs)
    assert 'consumer.py' in '\n'.join(p.text for p in check.paragraphs)
    shutil.move(str(temp), str(path))
finally:
    temp.unlink(missing_ok=True)
print('Inserted concise source tree; four embedded screenshots preserved.')
