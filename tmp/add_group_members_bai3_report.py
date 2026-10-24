"""Fill the group roster in the report cover page."""
from pathlib import Path
from tempfile import NamedTemporaryFile
import shutil

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

path = Path(r'D:\IOT-asm\bai3\bai3_1\Bao_cao_Bai3_Phan1_Nhom16.docx')
doc = Document(path)
label = next((p for p in doc.paragraphs if p.text.startswith('Thành viên nhóm')), None)
assert label is not None, 'Group roster line not found'
assert not any('Nguyễn Đăng Long' in cell.text for t in doc.tables for row in t.rows for cell in row.cells), 'Roster appears already inserted'

label.runs[0].text = 'Thành viên nhóm'
for run in label.runs[1:]:
    run.text = ''
label.runs[0].bold = True

members = [
    ('1', 'Nguyễn Đăng Long', '202400057', 'Thành viên'),
    ('2', 'Nguyễn Hoàng Gia', '202400040', 'Nhóm trưởng'),
    ('3', 'Nguyễn Trí Hiếu', '202416203', 'Thành viên'),
    ('4', 'Phạm Gia Linh', '202416262', 'Thành viên'),
]
table = doc.add_table(rows=1, cols=4)
table.autofit = False
table.alignment = 1
widths = [1.0, 7.4, 3.7, 4.2]
for col, width in zip(table.columns, widths):
    col.width = __import__('docx').shared.Cm(width)
headers = ['STT', 'Họ và tên', 'MSSV', 'Vai trò']
for cell, value in zip(table.rows[0].cells, headers):
    cell.text = value
for member in members:
    cells = table.add_row().cells
    for cell, value in zip(cells, member):
        cell.text = value

for ri, row in enumerate(table.rows):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))
    if ri == 0:
        tr_pr.append(OxmlElement('w:tblHeader'))
    for ci, cell in enumerate(row.cells):
        cell.width = __import__('docx').shared.Cm(widths[ci])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement('w:tcMar')
        for edge in ('top', 'left', 'bottom', 'right'):
            node = OxmlElement(f'w:{edge}')
            node.set(qn('w:w'), '55' if edge in ('top', 'bottom') else '75')
            node.set(qn('w:type'), 'dxa')
            mar.append(node)
        tcpr.append(mar)
        borders = OxmlElement('w:tcBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            node = OxmlElement(f'w:{edge}')
            node.set(qn('w:val'), 'single')
            node.set(qn('w:sz'), '4')
            node.set(qn('w:color'), 'D9D9D9')
            borders.append(node)
        tcpr.append(borders)
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'DCE6F1' if ri == 0 else ('F5F7FA' if ri % 2 == 0 else 'FFFFFF'))
        tcpr.append(shade)
        for paragraph in cell.paragraphs:
            paragraph.paragraph_format.space_after = 0
            if ci in (0, 2, 3):
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.name = 'Times New Roman'
                run.font.size = __import__('docx').shared.Pt(10)
                run.bold = ri == 0
                run._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')

table._tbl.getparent().remove(table._tbl)
label._p.addnext(table._tbl)

with NamedTemporaryFile(dir=path.parent, suffix='.docx', delete=False) as tmp:
    temp = Path(tmp.name)
try:
    doc.save(temp)
    check = Document(temp)
    assert len(check.inline_shapes) == 4, 'Embedded screenshots were changed'
    found = {(c.text) for t in check.tables for row in t.rows for c in row.cells}
    for _, name, student_id, role in members:
        assert name in found and student_id in found and role in found
    shutil.move(str(temp), str(path))
finally:
    temp.unlink(missing_ok=True)
print('Inserted four group members and verified all rows and four embedded screenshots.')
