from pathlib import Path
import json
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r'D:\IOT-asm\bai2_2')
WORK = Path(r'D:\IOT-asm\tmp\bai2_2_report')
OUTPUT = ROOT / 'Bao_cao_Bai_2_2_MQTT_Nguyen_Tri_Hieu_202416203.docx'
SRC = ROOT / 'src/main/java/vn/edu/iot'
sources = {n: (SRC / n).read_text(encoding='utf-8') for n in
           ('ClientConfig.java', 'Telemetry.java', 'MqttAssignment.java')}
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.6)
sec.left_margin, sec.right_margin = Cm(1.8), Cm(1.6)
sec.footer_distance = Cm(0.7)

for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption'):
    st = doc.styles[name]
    st.font.name = 'Times New Roman'
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.font.underline = False
    st.element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Times New Roman')
    st.element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Times New Roman')
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        st.element.get_or_add_rPr().rFonts.attrib.pop(qn('w:' + attr), None)
    st.paragraph_format.space_after = Pt(5)
    ppr = st.element.find(qn('w:pPr'))
    if ppr is not None:
        for border in list(ppr.findall(qn('w:pBdr'))):
            ppr.remove(border)
doc.styles['Normal'].font.size = Pt(12)
doc.styles['Normal'].paragraph_format.line_spacing = 1.08
doc.styles['Normal'].paragraph_format.widow_control = True
doc.styles['Title'].font.size = Pt(23)
doc.styles['Title'].font.bold = True
doc.styles['Subtitle'].font.size = Pt(15)
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
doc.styles['Heading 1'].paragraph_format.space_after = Pt(7)
doc.styles['Heading 2'].font.size = Pt(12)
doc.styles['Heading 2'].font.bold = True
doc.styles['Heading 2'].paragraph_format.space_before = Pt(9)
doc.styles['Heading 2'].paragraph_format.space_after = Pt(4)
doc.styles['Caption'].font.size = Pt(11)
doc.styles['Caption'].font.italic = True

code_style = doc.styles.add_style('Ma nguon', 1)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(9)
code_style.font.color.rgb = RGBColor(0, 0, 0)
for attr in ('ascii', 'hAnsi', 'eastAsia'):
    code_style.element.get_or_add_rPr().rFonts.set(qn('w:' + attr), 'Consolas')
code_style.paragraph_format.line_spacing = Pt(10.5)
code_style.paragraph_format.space_before = Pt(2)
code_style.paragraph_format.space_after = Pt(5)
code_style.paragraph_format.widow_control = True

def para(text, bold=False, keep=False):
    p = doc.add_paragraph()
    p.add_run(text).bold = bold
    p.paragraph_format.keep_with_next = keep
    return p

def heading(text, level=1, newpage=False):
    p = doc.add_heading(text, level)
    p.paragraph_format.page_break_before = newpage
    return p

def code(text):
    return doc.add_paragraph(text, style='Ma nguon')

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Cm(width)
    for cell, value in zip(t.rows[0].cells, headers):
        cell.text = value
        cell.paragraphs[0].runs[0].bold = True
        fill = OxmlElement('w:shd')
        fill.set(qn('w:fill'), 'E8EDF2')
        cell._tc.get_or_add_tcPr().append(fill)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for row in rows:
        for cell, value in zip(t.add_row().cells, row):
            cell.text = value
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + side)
        for k, v in {'val': 'single', 'sz': '5', 'color': 'D9D9D9'}.items():
            e.set(qn('w:' + k), v)
        borders.append(e)
    t._tbl.tblPr.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side in ('top', 'bottom', 'left', 'right'):
        e = OxmlElement('w:' + side)
        e.set(qn('w:w'), '85' if side in ('top', 'bottom') else '100')
        e.set(qn('w:type'), 'dxa')
        margins.append(e)
    t._tbl.tblPr.append(margins)
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(11)
    return t

manifest = []
def source_block(filename, start, end, title, explanation):
    heading(title, 2)
    p = para(f'Dòng {start} đến {end}. {explanation}', keep=True)
    lines = sources[filename].splitlines(keepends=True)
    raw = ''.join(lines[start - 1:end])
    display_lines = raw.rstrip('\n').split('\n')
    for i, line in enumerate(display_lines):
        p = code(line)
        p.paragraph_format.space_before = Pt(2 if i == 0 else 0)
        p.paragraph_format.space_after = Pt(5 if i == len(display_lines) - 1 else 0)
        if not line.strip():
            p.paragraph_format.line_spacing = Pt(4)
        p.paragraph_format.keep_together = True
        p.paragraph_format.keep_with_next = (i < 2 or
            (i >= len(display_lines) - 3 and i < len(display_lines) - 1))
    manifest.append({'file': filename, 'start': start, 'end': end, 'source': raw})

title = doc.add_paragraph('BÁO CÁO BÀI 2 PHẦN CÁ NHÂN', style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(12)
title.paragraph_format.space_after = Pt(8)
p = doc.add_paragraph('Trao đổi dữ liệu MQTT bằng Java', style='Subtitle')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(14)
for text in ('Sinh viên: Nguyễn Trí Hiếu', 'MSSV: 202416203'):
    p = para(text, True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
para('')
heading('1 Nội dung thực hiện')
para('Chương trình Java vừa gửi dữ liệu cảm biến mô phỏng lên broker, vừa nhận bản tin MQTT và hiển thị đủ các trường JSON. Báo cáo trình bày toàn bộ ba tệp mã nguồn chính, giải thích từng khối và hướng dẫn chụp ảnh minh chứng khi chạy thực tế.')
table(('Thông số', 'Giá trị sử dụng'), [
    ('Môi trường', 'JDK 17 trở lên; Eclipse Paho 1.2.5; JSON-java 20240303'),
    ('Broker mặc định', 'tcp://broker.hivemq.com:1883'),
    ('Topic gửi dữ liệu', 'it4735e/202416203/telemetry'),
    ('Topic đăng ký nhận', 'it4735e/202416203/#'),
    ('Topic trạng thái', 'it4735e/202416203/status'),
    ('Gửi dữ liệu', 'QoS 1; packet_no bắt đầu từ 126; thời gian chờ mặc định 5 giây giữa các lượt gửi')
], (4.1, 13.5))
heading('Vai trò của ba tệp', 2)
para('ClientConfig.java kiểm tra tham số và tạo topic theo MSSV. Telemetry.java tạo dữ liệu mô phỏng, chuyển đổi JSON và định dạng kết quả. MqttAssignment.java điều khiển kết nối, gửi nhận, phục hồi kết nối và dừng chương trình.')
heading('Luồng hoạt động', 2)
para('Đọc tham số → kết nối broker → đăng ký topic → gửi trạng thái online → gửi telemetry định kỳ. Mỗi bản tin nhận được sẽ được phân loại thành trạng thái hoặc dữ liệu JSON để hiển thị.')
para('Dữ liệu gồm 7 trường: id, name, packet_no, temperature, humidity, tds và pH. Với QoS 1, bản tin nhận có thể trùng; packet_no được đặt lại về 126 mỗi khi chạy chương trình mới.')

heading('2 Cấu hình kết nối', newpage=True)
para('Tệp ClientConfig.java', True)
source_block('ClientConfig.java', 1, 24, 'Khai báo và kiểm tra cấu hình',
    'Record lưu MSSV, broker và khoảng chờ gửi. Hàm khởi tạo yêu cầu MSSV hợp lệ, khoảng chờ dương và địa chỉ broker đúng dạng tcp://host:port; cấu hình sai sẽ phát sinh lỗi để báo cho người dùng.')
source_block('ClientConfig.java', 25, 53, 'Đọc tham số dòng lệnh',
    'parse đọc --mssv, --broker và --interval. Broker mặc định là HiveMQ, khoảng chờ là 5 giây. Chương trình báo lỗi nếu gặp tùy chọn lạ, thiếu giá trị hoặc khoảng chờ không phải số nguyên; MSSV phải được cung cấp.')
source_block('ClientConfig.java', 54, 66, 'Tạo topic riêng theo MSSV',
    'Ba hàm ghép MSSV vào topic gửi telemetry, topic đăng ký có ký tự # và topic trạng thái. Khi dùng MSSV 202416203, các topic trùng với bảng thông số ở phần đầu.')

heading('3 Dữ liệu cảm biến')
para('Tệp Telemetry.java', True)
source_block('Telemetry.java', 1, 18, 'Khai báo bản tin và tạo dữ liệu mô phỏng',
    'Record chứa đủ 7 trường. sample tạo nhiệt độ và pH làm tròn một chữ số thập phân, độ ẩm và TDS ngẫu nhiên; packetNo lấy từ bộ đếm gửi. Đây là dữ liệu mô phỏng, chưa đọc từ cảm biến thật.')
source_block('Telemetry.java', 19, 30, 'Chuyển dữ liệu thành JSON',
    'toJson dùng JSONObject để tạo chuỗi gửi đi. Tên khóa packet_no dùng dấu gạch dưới và pH giữ đúng chữ hoa chữ thường theo yêu cầu.')
source_block('Telemetry.java', 31, 42, 'Đọc và kiểm tra JSON nhận được',
    'fromJson đọc đối tượng JSON, kiểm tra phần nội dung thừa phía sau rồi lấy đủ các trường. Các hàm get báo lỗi khi thiếu trường hoặc giá trị không thể đọc theo kiểu yêu cầu; lỗi sẽ được xử lý ở lớp MQTT.')
source_block('Telemetry.java', 43, 50, 'Định dạng kết quả để hiển thị',
    'display tạo hai dòng kết quả chứa đủ thông tin bản tin. Locale.ROOT giữ cách định dạng số nhất quán, các giá trị cảm biến được in với một chữ số thập phân.')

heading('4 Chương trình gửi nhận MQTT')
para('Tệp MqttAssignment.java', True)
source_block('MqttAssignment.java', 1, 32, 'Thư viện và trạng thái chương trình',
    'Các thư viện hỗ trợ MQTT, JSON, UTF-8 và xử lý đồng thời. QoS đặt bằng 1, thời gian chờ xác nhận tối đa 15 giây. worker thực hiện công việc theo lịch trên một luồng; ready và closing kiểm soát việc gửi, stopped giữ chương trình chờ, bộ đếm bắt đầu từ 126.')
source_block('MqttAssignment.java', 33, 51, 'Tạo client và xử lý kết nối thành công',
    'UUID tạo client ID mới cho mỗi lần chạy. Callback connectComplete gọi prepareSession trên worker sau cả kết nối đầu tiên và kết nối lại. Việc chờ xác nhận đăng ký hoặc gửi được chuyển sang worker để không chặn luồng callback của Paho.')
source_block('MqttAssignment.java', 52, 74, 'Xử lý mất kết nối và nhận bản tin',
    'connectionLost đặt ready về false để tạm ngừng gửi. messageArrived giải mã payload UTF-8 rồi gọi formatReceived. deliveryComplete không cần xử lý thêm vì hàm gửi tự chờ xác nhận trên worker; cuối hàm khởi tạo in thông tin client và topic.')
source_block('MqttAssignment.java', 75, 86, 'Thiết lập MQTT và Last Will',
    'Dùng MQTT 3.1.1, phiên sạch, thời gian kết nối 10 giây và keep alive 15 giây. Tự kết nối lại được bật. Last Will là chuỗi offline trên topic trạng thái, QoS 1 và retained để broker lưu trạng thái gần nhất khi kết nối mất bất thường.')
source_block('MqttAssignment.java', 87, 106, 'Kết nối lần đầu và lên lịch gửi',
    'Nếu lần kết nối đầu thất bại, chương trình chờ 5 giây rồi thử lại. Khi kết nối thành công, scheduleWithFixedDelay lên lịch gửi: sau khi một lượt xử lý kết thúc, chờ thêm khoảng thời gian đã cấu hình rồi chạy lượt tiếp theo. stopped.await giữ chương trình hoạt động đến lúc đóng.')
source_block('MqttAssignment.java', 107, 131, 'Chuẩn bị phiên sau mỗi lần kết nối',
    'prepareSession đăng ký topic có # với QoS 1 và chờ broker xác nhận. Sau đó gửi online với retained rồi bật ready. Vì cleanSession bằng true nên phải đăng ký lại sau mỗi lần kết nối; nếu chuẩn bị phiên lỗi, chương trình thử lại sau 2 giây.')
source_block('MqttAssignment.java', 132, 149, 'Gửi dữ liệu theo lịch và tăng bộ đếm',
    'Chỉ gửi khi client đã sẵn sàng và còn kết nối. Dữ liệu được tạo bằng Telemetry.sample, đổi sang JSON rồi gửi QoS 1, không retained. Bộ đếm tăng ngay khi lệnh publish được chấp nhận, trước khi chờ xác nhận; lỗi xác nhận được in ra và chương trình vẫn tiếp tục.')
source_block('MqttAssignment.java', 150, 161, 'Phân loại trạng thái và xử lý JSON lỗi',
    'Bản tin trên topic status được in như văn bản. Các topic còn lại được đọc thành Telemetry và hiển thị đủ trường. JSONException được bắt tại đây để bản tin sai không làm lỗi lan ra callback; bản tin hợp lệ tiếp theo vẫn được xử lý.')
source_block('MqttAssignment.java', 162, 193, 'Đóng kết nối và giải phóng tài nguyên',
    'compareAndSet bảo đảm việc đóng chỉ thực hiện một lần. Chương trình dừng worker, chủ động gửi offline retained khi còn kết nối rồi ngắt MQTT. Khối finally dọn kết nối, đóng client và mở chốt stopped. Khi thoát bình thường, offline được chương trình tự gửi vì Last Will dành cho mất kết nối bất thường.')
source_block('MqttAssignment.java', 194, 228, 'Điểm bắt đầu và xử lý khi thoát',
    'main hỗ trợ --help, đọc cấu hình và báo lỗi tham số. Shutdown hook gọi close khi JVM kết thúc, ví dụ Ctrl+C; try-with-resources và finally bảo đảm dọn tài nguyên. Nếu luồng bị ngắt, chương trình đặt lại cờ ngắt; lỗi tạo client được báo riêng.')
source_block('MqttAssignment.java', 229, 234, 'Hiển thị cách dùng',
    'usage in mẫu lệnh gồm MSSV bắt buộc, khoảng chờ gửi và địa chỉ broker tùy chọn. Khi chạy bài này, thay YOUR_MSSV bằng 202416203.')

def screenshot(num, title, steps, expected, height=122):
    heading(f'Hình {num} {title}', 2)
    p = para('Thực hiện: ' + steps)
    p.paragraph_format.keep_with_next = True
    p = para('Cần thấy trong ảnh: ' + expected)
    p.paragraph_format.keep_with_next = True
    p = para(f'[CHÈN ẢNH CHỤP MÀN HÌNH SỐ {num} TẠI ĐÂY]')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(height / 2)
    p.paragraph_format.space_after = Pt(height / 2)
    p.runs[0].font.size = Pt(11)
    p.runs[0].font.color.rgb = RGBColor.from_string('777777')

heading('5 Thực hành và ảnh minh chứng', newpage=True)
para('Chạy bằng JDK 17 trở lên. Mở PowerShell tại thư mục bai2_2 và dùng lệnh sau; dừng bằng Ctrl+C. Nếu PowerShell chặn script, dùng lệnh javac và java bên dưới với hai thư viện đã có trong lib.')
code('cd D:\\IOT-asm\\bai2_2\n.\\run.ps1 -Mssv 202416203')
code('javac --release 17 -encoding UTF-8 -cp "lib/*" -d out src/main/java/vn/edu/iot/*.java\njava -cp "out;lib/*" vn.edu.iot.MqttAssignment --mssv 202416203')
para('Trong MQTTX hoặc MQTTBox, kết nối TCP tới broker.hivemq.com, cổng 1883, không nhập tài khoản; dùng client ID khác Java và đăng ký it4735e/202416203/# với QoS 1. Chụp kết quả thực tế vào các vị trí dưới đây; thêm nhận xét nếu có khác biệt.')
screenshot(1, 'Java kết nối và đăng ký topic',
    'Khởi chạy Java và chờ kết nối thành công.',
    'Client ID, Connected to, Subscribed: it4735e/202416203/# và thông báo gửi online.', 110)
screenshot(2, 'Công cụ nhận dữ liệu do Java gửi',
    'Giữ MQTTX hoặc MQTTBox đang đăng ký topic và đợi ít nhất ba bản tin telemetry.',
    'Topic chứa 202416203, QoS 1, đủ 7 trường và packet_no tăng qua các bản tin mới, ví dụ 126, 127, 128.', 110)

heading('Ảnh kiểm tra dữ liệu nhận và xử lý lỗi', newpage=True)
para('Gửi JSON mẫu sau từ công cụ tới it4735e/202416203/telemetry với QoS 1. Để name bằng mqttx hoặc mqttbox giúp phân biệt với bản tin Java tự gửi.')
code('{"id":11,"name":"mqttx","packet_no":1,"temperature":25,\n "humidity":55,"tds":900,"pH":6.5}')
screenshot(3, 'Java nhận JSON từ công cụ',
    'Publish JSON mẫu, sau đó xem cửa sổ Java; có thể chụp hai cửa sổ cạnh nhau.',
    'Topic gửi của công cụ và dòng Received của Java có name=mqttx, packet_no=1 cùng đủ 7 trường.', 140)
screenshot(4, 'Bản tin sai JSON và phục hồi',
    'Chuyển payload sang văn bản thô, gửi hello vào topic telemetry; sau đó gửi lại JSON mẫu hợp lệ.',
    'JSON error có continuing, tiếp theo là bản tin hợp lệ được đọc hoặc telemetry tiếp tục gửi mà không phải khởi động lại Java.', 140)

heading('Ảnh kiểm tra các chức năng mở rộng', newpage=True)
para('Hai nội dung dưới đây minh họa Last Will, retained và tự kết nối lại đã được cài đặt trong mã nguồn. Dùng một phiên Java cho MSSV 202416203 để trạng thái dễ theo dõi.')
screenshot(5, 'Last Will và trạng thái retained',
    'Giữ công cụ kết nối. Khi Java đang online, kết thúc bất thường đúng tiến trình Java trong Task Manager; đợi broker báo offline. Kết nối lại công cụ và đăng ký topic để kiểm tra trạng thái lưu.',
    'Topic it4735e/202416203/status đổi từ online sang offline; một subscriber mới nhận ngay giá trị offline đã lưu. Có thể chèn hai ảnh. Ctrl+C là thoát bình thường nên không dùng để chứng minh Last Will.', 140)
screenshot(6, 'Tự kết nối lại và đăng ký lại',
    'Khởi chạy lại Java, tạm ngắt mạng của máy chạy Java rồi nối lại và đợi chương trình phục hồi.',
    'Connection lost, Reconnected, Subscribed, online và telemetry tiếp tục xuất hiện sau khi có mạng.', 140)
para('Nhận xét sau khi thực hành: ................................................................................................\n...........................................................................................................................................')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Trang ').font.size = Pt(10)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Báo cáo bài 2 phần cá nhân Trao đổi dữ liệu MQTT bằng Java'
doc.core_properties.author = 'Nguyễn Trí Hiếu'
doc.core_properties.subject = 'MSSV 202416203 Mã nguồn và hướng dẫn minh chứng thực hành'
doc.core_properties.keywords = 'MQTT, Java, 202416203'
doc.save(OUTPUT)
WORK.mkdir(parents=True, exist_ok=True)
(WORK / 'source_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
for filename, source in sources.items():
    pieces = [b for b in manifest if b['file'] == filename]
    assert ''.join(b['source'] for b in pieces) == source, filename
    assert pieces[0]['start'] == 1
    assert pieces[-1]['end'] == len(source.splitlines())
print(OUTPUT)
print('Verified all 350 source lines across three files; six screenshot positions.')
