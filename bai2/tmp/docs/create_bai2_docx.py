from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

root = Path('D:/IOT-asm/bai2')
source = (root / 'main.py').read_text(encoding='utf-8')
doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)

for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Arial')
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing = 1.12
doc.styles['Normal'].paragraph_format.space_after = Pt(7)
doc.styles['Title'].font.size = Pt(19)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(7)
doc.styles['Subtitle'].font.size = Pt(11)
doc.styles['Subtitle'].paragraph_format.space_after = Pt(12)
for name, size in [('Heading 1', 13), ('Heading 2', 11.5)]:
    style = doc.styles[name]
    style.font.size = Pt(size)
    style.font.bold = True
    style.paragraph_format.space_before = Pt(8)
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.keep_with_next = True
code_style = doc.styles.add_style('Ma nguon', WD_STYLE_TYPE.PARAGRAPH)
code_style.font.name = 'Consolas'
code_style.font.size = Pt(8.5)
code_style.font.color.rgb = RGBColor(0, 0, 0)
code_style.paragraph_format.space_after = Pt(0)
code_style.paragraph_format.line_spacing = Pt(10.5)
code_style.paragraph_format.keep_together = True
code_style.paragraph_format.widow_control = False

def p(text, style=None):
    return doc.add_paragraph(text, style)
def h(text, level=1):
    return doc.add_heading(text, level)
def code(text):
    for line in text.strip('\n').splitlines():
        p(line, 'Ma nguon')
def image_slot(label):
    para = p(label)
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(60)
    para.paragraph_format.space_after = Pt(75)
    para.runs[0].italic = True
    para.runs[0].font.color.rgb = RGBColor(110, 110, 110)

p('Báo cáo bài tập cá nhân', 'Title')
p('Trao đổi dữ liệu sử dụng HTTP API với ThingSpeak', 'Subtitle')
p('Họ và tên: [Điền họ và tên]')
p('Mã số sinh viên: [Điền mã số sinh viên]')
p('Chương trình Python gửi nhiệt độ và độ ẩm lên ThingSpeak bằng GET và POST JSON, sau đó lấy hai bản ghi mới nhất, phân tích JSON và hiển thị dữ liệu. field1 biểu diễn nhiệt độ (°C); field2 biểu diễn độ ẩm (%).')
h('1 Mã nguồn gửi dữ liệu')
p('File main.py dùng thư viện chuẩn Python. Các giá trị mẫu là nhiệt độ 20°C và độ ẩm 33%. API key và Channel ID được cấu hình theo channel sử dụng.')
code(source[source.index('WRITE_API_KEY ='):source.index('def read_data():')].strip())
p('Hàm send_data(method) dùng urlencode() để đóng gói dữ liệu trong URL khi gửi GET. Khi gửi POST, json.dumps() tạo body JSON và đặt Content-Type: application/json. Hàm đọc entry_id trả về; 0 là cập nhật thất bại, khác 0 là thành công.')

doc.add_page_break()
h('2 Mã nguồn đọc dữ liệu và điều phối')
h('Đọc dữ liệu bằng GET', 2)
code(source[source.index('def read_data():'):source.index('def main():')].strip())
p('Hàm read_data() gửi GET đến feeds.json với results=2. json.loads() phân tích phản hồi, sau đó vòng lặp lấy field1 và field2 trong mảng feeds và in nhiệt độ, độ ẩm. Giá trị null được hiển thị là không có dữ liệu.')
h('Luồng chạy chương trình', 2)
code(source[source.index('def main():'):].strip())
p('Hàm main() gửi GET, chờ 16 giây, gửi POST JSON, rồi đọc hai bản ghi mới nhất. Khoảng chờ đáp ứng yêu cầu tối thiểu 15 giây giữa hai lần cập nhật. Chương trình hiển thị lỗi nếu request thất bại.')

doc.add_page_break()
h('3 Kết quả thực hiện')
h('Kết quả chạy chương trình Python', 2)
code('python main.py')
p('Kết quả GET: entry_id = [Điền kết quả]. Kết quả POST: entry_id = [Điền kết quả]. Hai bản ghi đọc về: [Điền nhiệt độ và độ ẩm thực tế].')
image_slot('[Chèn ảnh terminal khi chạy main.py tại đây]')
p('Hình 1 Chương trình gửi GET và POST rồi hiển thị nhiệt độ và độ ẩm của hai bản ghi mới nhất', 'Caption')
h('Postman gửi dữ liệu bằng GET', 2)
code('GET https://api.thingspeak.com/update\nQuery: api_key=<WRITE_API_KEY>, field1=20, field2=33')
p('Phản hồi: HTTP [Điền mã trạng thái], entry_id = [Điền kết quả].')
image_slot('[Chèn ảnh Postman gửi GET và phản hồi tại đây]')
p('Hình 2 Request GET có field1 và field2 trong query string cùng phản hồi', 'Caption')

doc.add_page_break()
h('Kết quả thực hiện bằng Postman')
h('Postman gửi dữ liệu bằng POST JSON', 2)
code('POST https://api.thingspeak.com/update?api_key=<WRITE_API_KEY>\nContent-Type: application/json\n{"field1": 20, "field2": 33}')
p('Phản hồi: HTTP [Điền mã trạng thái], entry_id = [Điền kết quả].')
image_slot('[Chèn ảnh Postman gửi POST JSON và phản hồi tại đây]')
p('Hình 3 Request POST có header application/json và body chứa hai trường dữ liệu cùng phản hồi', 'Caption')
h('Postman đọc dữ liệu bằng GET', 2)
code('GET https://api.thingspeak.com/channels/<CHANNEL_ID>/feeds.json?results=2')
p('Phản hồi: HTTP [Điền mã trạng thái]. Giá trị field1 và field2 trong hai phần tử feeds: [Điền kết quả].')
image_slot('[Chèn ảnh Postman đọc dữ liệu và JSON phản hồi tại đây]')
p('Hình 4 Phản hồi JSON chứa hai bản ghi mới nhất trong mảng feeds', 'Caption')

props = doc.core_properties
props.title = 'Báo cáo bài tập cá nhân HTTP API với ThingSpeak'
props.subject = 'Phần 2 Bài tập về nhà cá nhân'
props.author = ''
props.last_modified_by = ''
props.comments = ''
output = root / 'output/bao-cao.docx'
doc.save(output)
print(output)
