from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\IOT-asm\bai1")
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
BLACK = RGBColor(0, 0, 0)


def font(style, name, size):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{key}"), name)


def document(title, source=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.6 if source else 0.65)
    section.bottom_margin = Inches(0.6 if source else 0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = doc.styles
    for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2"):
        font(styles[name], "Arial", 11)
        props = styles[name].element.get_or_add_pPr()
        border = props.find(qn("w:pBdr"))
        if border is not None:
            props.remove(border)
    normal = styles["Normal"].paragraph_format
    normal.line_spacing = 1.08
    normal.space_after = Pt(5)
    normal.widow_control = True
    styles["Title"].font.size = Pt(19)
    styles["Title"].font.bold = True
    styles["Title"].paragraph_format.space_after = Pt(5)
    styles["Title"].paragraph_format.line_spacing = 1.02
    styles["Title"].paragraph_format.keep_with_next = True
    styles["Subtitle"].font.size = Pt(10)
    styles["Subtitle"].paragraph_format.space_after = Pt(9)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.size = Pt(12)
        styles[name].font.bold = True
        styles[name].paragraph_format.space_before = Pt(8)
        styles[name].paragraph_format.space_after = Pt(4)
        styles[name].paragraph_format.keep_with_next = True

    for name, face, size, spacing in (
        ("Source Code", "Consolas", 10, 11.5),
        ("Command", "Consolas", 10, 12),
        ("Small", "Arial", 9.5, 11.5),
    ):
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        font(style, face, size)
        pf = style.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0 if name == "Source Code" else 4)
        pf.line_spacing = Pt(spacing)
        pf.keep_together = True
        pf.widow_control = False if name == "Source Code" else True

    doc.core_properties.title = title
    doc.core_properties.subject = "Bài 1 Phần 1 Trao đổi dữ liệu sử dụng HTTP API"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "HTTP, GET, JSON, Python, nhiệt độ, độ ẩm"

    if source:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Trang ")
        run.font.name = "Arial"
        run.font.size = Pt(9)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)
    return doc


def p(doc, text, style=None):
    return doc.add_paragraph(text, style)


def heading(doc, text):
    return doc.add_paragraph(text, "Heading 1")


report = document("Báo cáo trao đổi dữ liệu cảm biến qua HTTP")
p(report, "Báo cáo trao đổi dữ liệu\ncảm biến qua HTTP", "Title")
p(report, "Bài 1 Phần 1", "Subtitle")
p(report, "Nhóm: .....................................    Lớp: .....................................", "Small")
p(report, "Thành viên: ........................................................................................", "Small")

heading(report, "1 Mục tiêu")
p(report, "Xây dựng server cung cấp nhiệt độ và độ ẩm; client gửi HTTP GET để nhận dữ liệu JSON và hiển thị kết quả trong terminal. Chương trình gồm server.py và client.py, sử dụng thư viện chuẩn của Python 3.9 trở lên.")

heading(report, "2 Thiết kế và giải thích chương trình")
p(report, "Server lắng nghe tại 127.0.0.1:8000 và xử lý GET /api/sensor. Mỗi request tạo dữ liệu giả lập mới: nhiệt độ từ 20 đến 35 °C, làm tròn một chữ số thập phân; độ ẩm từ 40 đến 80%, là số nguyên. Dữ liệu được sinh ngẫu nhiên, không đọc từ cảm biến vật lý; các lần gọi có thể trả giá trị trùng nhau.")
p(report, "Hàm generate_sensor_data() dùng random.uniform() và random.randint() để tạo dữ liệu. do_GET() chuyển dữ liệu thành JSON bằng json.dumps(), gửi HTTP 200 cùng Content-Type: application/json; charset=utf-8. Cache-Control: no-store yêu cầu không lưu phản hồi vào bộ nhớ đệm; đường dẫn sai trả HTTP 404.")
p(report, "Client dùng Request(..., method=\"GET\") và urlopen() với thời gian chờ 10 giây. json.loads() phân tích phản hồi; client kiểm tra giá trị là số hữu hạn, độ ẩm trong khoảng 0 đến 100%, rồi in nhiệt độ và độ ẩm nhận được. Client báo lỗi khi HTTP, kết nối hoặc dữ liệu không hợp lệ.")

heading(report, "3 Cách chạy")
p(report, "Lưu mã nguồn trong tài liệu đi kèm thành hai file .py. Mở hai terminal trong cùng thư mục; chạy server trước, sau đó chạy client.")
p(report, "Terminal 1: python server.py\nTerminal 2: python client.py", "Command")
p(report, "Chạy lại client để lấy dữ liệu mới. Nhấn Ctrl+C tại terminal server để dừng. Nếu cổng 8000 đã được dùng, chạy server.py --port 8080 và client.py --url http://127.0.0.1:8080/api/sensor.", "Small")

heading(report, "4 Kết quả và ảnh minh họa")
p(report, "Kiểm tra 20 request GET cho thấy server trả JSON hợp lệ và dữ liệu nằm trong các khoảng đã cấu hình. Client hiển thị đúng dữ liệu nhận được; đường dẫn sai trả 404 và client báo lỗi. Ví dụ kết quả: nhiệt độ 27.3 °C, độ ẩm 73%.")
p(report, "Hình 1 Ảnh terminal server đang hoạt động và ghi nhận request GET", "Small")
p(report, "[Chèn ảnh chụp màn hình server tại đây]", "Small")
p(report, "Hình 2 Ảnh terminal client hiển thị HTTP 200 JSON nhiệt độ và độ ẩm", "Small")
p(report, "[Chèn ảnh chụp màn hình client tại đây]", "Small")
report_path = OUT / "Bao_cao_HTTP_cam_bien.docx"
report.save(report_path)


source = document("Mã nguồn chương trình cảm biến HTTP", source=True)
p(source, "Mã nguồn chương trình cảm biến HTTP", "Title")
p(source, "Bài 1 Phần 1", "Subtitle")
p(source, "Hai chương trình Python dưới đây trao đổi dữ liệu cảm biến giả lập qua HTTP GET. Sao chép từng phần thành file có tên tương ứng và giữ nguyên thụt lề.", "Small")

for index, name in enumerate(("server.py", "client.py"), start=1):
    heading(source, f"{index} Mã nguồn {'server' if index == 1 else 'client'}")
    p(source, f"Tên file: {name}", "Small")
    text = (ROOT / name).read_text(encoding="utf-8")
    for line in text.splitlines():
        p(source, line, "Source Code")

source_path = OUT / "Ma_nguon_server_client.docx"
source.save(source_path)

# Doi chieu toan bo ma nguon sau khi luu de bao dam khong mat dong/thut le.
saved = Document(source_path)
code_groups = []
current = []
for paragraph in saved.paragraphs:
    if paragraph.style.name == "Source Code":
        current.append(paragraph.text)
    elif current:
        code_groups.append(current)
        current = []
if current:
    code_groups.append(current)
assert len(code_groups) == 2
for name, lines in zip(("server.py", "client.py"), code_groups):
    expected = (ROOT / name).read_text(encoding="utf-8").splitlines()
    assert lines == expected, f"Ma nguon trong Word khong khop {name}"
    compile("\n".join(lines) + "\n", name, "exec")
print(report_path)
print(source_path)
print("PASS: source document contains both complete Python files with exact indentation")
